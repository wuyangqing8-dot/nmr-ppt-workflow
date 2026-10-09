"""Extract manual labels from the reference without molecule-specific lookup tables.

PPT text boxes carry no chemical atom/peak links. Positional matches are explicitly
provisional and never counted as verified chemical assignments.
"""
from collections import Counter
from pathlib import Path
from pptx import Presentation
from app.annotation import atom_to_slide
from app.storage import save_json

def compare(task,template,out):
    p=Presentation(template);slide=p.slides[0];seen=Counter();struct=[];spectral=[]
    for sh in slide.shapes:
        if not sh.has_text_frame:continue
        text=sh.text.strip()
        box=[v/914400 for v in [sh.left,sh.top,sh.width,sh.height]]
        row={'label':text,'box':box}
        if text.isalpha() and text.islower():
            (struct if seen[text]==0 else spectral).append(row);seen[text]+=1
        elif '+' in text and all(part.isalpha() and part.islower() for part in text.split('+')):spectral.append(row)
    original_box=task['layout']['structure_box']
    locations={a['id']:atom_to_slide(a['xy'],task['structure']['bbox'],original_box) for a in task['structure']['atoms']}
    rows=[]
    for row in struct:
        x,y,w,h=row['box'];cx=x+w/2;cy=y+h/2
        nearest=sorted(locations,key=lambda aid:(locations[aid][0]-cx)**2+(locations[aid][1]-cy)**2)[:3]
        candidates=[e for e in task['structure']['environments'] if nearest[0] in e['atom_ids']]
        peaks=[s for s in spectral if row['label'] in s['label'].split('+')]
        rows.append({'reference_label':row['label'],'nearest_cdxml_atoms':nearest,
            'candidate_environments':[e['label'] for e in candidates],
            'predicted_ppm':[e.get('predicted_ppm') for e in candidates],
            'suggested_experimental_ppm':[e.get('experimental_ppm') for e in candidates],
            'reference_peak_textboxes':peaks,'verified_correct':False,
            'reason':'参考文本框无原子/ppm 语义映射；邻近原子仅用于人工对照，不能作为已验证真值'})
    result={'reference_structure_labels':len(struct),'reference_spectrum_labels':len(spectral),
        'overlap_labels':[s['label'] for s in spectral if '+' in s['label']],
        'verified_correct_count':0,'correct_count_status':'无法由无语义链接的 PPT 文本框独立核实',
        'needs_review':len(task['structure']['environments']),'rows':rows,
        'note':'模板标签首次出现按结构、后续按谱图归类；需人工检查提取顺序。自动归属程序不调用此比较器。'}
    save_json(Path(out)/'reference_comparison.json',result)
    lines=['# 自动归属与人工参考的对照','',
        f'参考提取到结构标签 {len(struct)} 个、谱图标签 {len(spectral)} 个，组合标签：'+','.join(result['overlap_labels']),
        f'已独立核实正确：0；自动候选待审核：{result["needs_review"]}。',
        '0 表示尚未取得可审计的人工原子到峰映射，不表示所有建议错误。',
        '参考 PPT 的文本框没有原子 ID 或 ppm 属性，最近原子可能误配。下表仅供逐项审核；不把相似位置当作成功归属。',
        '', '|人工标签|近邻 CDXML 原子|自动环境候选|真实预测 ppm|实验候选 ppm|','|---|---|---|---|---|']
    for row in rows:lines.append('|'+ '|'.join(str(row[k]) for k in ['reference_label','nearest_cdxml_atoms','candidate_environments','predicted_ppm','suggested_experimental_ppm'])+'|')
    (Path(out)/'人工参考对照.md').write_text('\n'.join(lines),encoding='utf-8')
    return result
