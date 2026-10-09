import copy
from pathlib import Path
import numpy as np
import pytest
from app.annotation import ppm_to_x
from app.assignment import auto_assign,make_signals
from app.prediction_parser import attach_prediction,attach_csv
from app.structure_parser import parse_cdxml
from app.validation import validate
from app.storage import read_json
from app.nmr_import import read_xy,spectrum_from_export

def env(i,ppm,h):
    return dict(id=f'E{i}',label=chr(96+i),predicted_ppm=ppm,prediction_error=.2,hydrogen_count=h,
        environment='CH',reason='',status='尚未解决',confidence=0.,atom_ids=['1'],prediction_records=[],positions={})

def test_ppm_axis():
    assert ppm_to_x(8,8,0,2,8)==2
    assert ppm_to_x(0,8,0,2,8)==10
    assert ppm_to_x(5.3,8,0,2,8)==pytest.approx(4.7)
    with pytest.raises(ValueError):ppm_to_x(1,0,8,0,8)

def test_real_overlap_aggregate_integral():
    es=[env(1,2.00,2),env(2,2.02,3)]
    ss=[dict(id='P1',ppm=2.01,integral=5,multiplicity='',j_hz=[])]
    auto_assign(es,ss,integral_scale=1)
    assert all(e['signal_id']=='P1' and e['status']=='待确认' for e in es)
    assert es[0]['overlap_labels']==['a','b']
    assert ss[0]['integral']==5 and es[0]['integral_mismatch']==0

def test_no_forced_remote_assignment():
    es=[env(1,7.8,1)];ss=[dict(id='P1',ppm=.9,integral=None)]
    auto_assign(es,ss)
    assert es[0]['experimental_ppm'] is None and es[0]['status']=='尚未解决'

def test_confirmed_assignment_is_a_global_constraint():
    es=[env(1,2.,2),env(2,2.01,2)]
    es[0].update(status='人工确认',signal_id='P2',experimental_ppm=2.3,integral=2)
    ss=[dict(id='P1',ppm=2.,integral=2),dict(id='P2',ppm=2.3,integral=2)]
    auto_assign(es,ss,integral_scale=1)
    assert es[0]['signal_id']=='P2' and es[0]['status']=='人工确认'
    assert es[1]['signal_id']=='P1'
    with pytest.raises(ValueError,match='冲突'):
        auto_assign(es,ss,[dict(signal_id='P2',status='人工确认')])

def test_xy_sort_and_validation(tmp_path):
    p=tmp_path/'xy.txt';p.write_text('ppm,intensity\n0,0\n1,5\n2,1')
    sp=read_xy(p);assert sp['ppm']==[2.,1.,0.];assert sp['intensity']==[1.,5.,0.]
    p.write_text('ppm,intensity\n0,0\n2,5\n1,1')
    with pytest.raises(ValueError):read_xy(p)

def test_unmapped_prediction_is_not_atom_prediction(tmp_path):
    p=tmp_path/'pred.csv';p.write_text('ppm,h_count\n7.8,1')
    s={'atoms':[{'id':'1'}],'environments':[env(1,None,1)]}
    result=attach_csv(s,p)
    assert result['mapped_environments']==0 and s['environments'][0]['predicted_ppm'] is None

ROOT=Path(__file__).resolve().parents[1]
TASK=ROOT/'outputs/OP-3TH-OP-auto/task.json'

@pytest.mark.skipif(not TASK.exists(),reason='real case generated separately')
def test_real_structure_prediction_and_integral_provenance():
    t=read_json(TASK);s,m=parse_cdxml(t['assets']['cdxml'])
    assert len(s['atoms'])==m.GetNumAtoms()==195
    assert len(s['bonds'])==m.GetNumBonds()==209
    assert sum(e['hydrogen_count'] for e in t['structure']['environments'])==266
    assert t['prediction']['structure_consistent'] is True
    assert t['prediction']['records']==145
    assert validate(t)
    broken=copy.deepcopy(t);broken['spectrum']['integrals'][0]['value']+=1
    with pytest.raises(ValueError,match='原始积分'):validate(broken)

@pytest.mark.skipif(not TASK.exists(),reason='real case generated separately')
def test_export_has_no_duplicate_structure_and_editable_labels():
    from pptx import Presentation
    p=Presentation(ROOT/'outputs/OP-3TH-OP-auto/assignment_draft.pptx')
    pics=[s for s in p.slides[0].shapes if s.shape_type==13]
    assert len(pics)==2
    assert any(s.has_text_frame and s.name.startswith('PeakLabel_') for s in p.slides[0].shapes)
    assert any(s.has_text_frame and s.name.startswith('StructureLabel_') for s in p.slides[0].shapes)
    from collections import Counter
    counts=Counter(s.name.split('_')[1] for s in p.slides[0].shapes if s.name.startswith('StructureLabel_'))
    task=read_json(TASK)
    assert counts=={e['id']:1 for e in task['structure']['environments']}
    assert next(s for s in p.slides[0].shapes if s.name=='MolecularFormula').text==task['structure']['formula']

def test_reference_render_is_identical():
    from PIL import Image,ImageChops
    a=ROOT/'outputs/reference/幻灯片1.PNG';b=ROOT/'outputs/OP-3TH-OP-reference/preview.png'
    if not a.exists() or not b.exists():pytest.skip('reference render not available')
    assert ImageChops.difference(Image.open(a).convert('RGB'),Image.open(b).convert('RGB')).getbbox() is None
